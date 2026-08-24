"""Recalculate the generated workbook and assert the suitability outcomes.

Runs the real Excel formulas through the `formulas` engine, so a syntax error or a
broken lookup range fails here rather than in front of a user.

    python verify_workbook.py              # every case, every regulation tab
    python verify_workbook.py "FCA v15"    # only cases that run on that tab

There is one calculator tab per regulation. Cell addresses are resolved by scanning
row labels rather than hardcoded, so a tab can be re-laid-out without silently
invalidating these checks.
"""

import sys

import formulas
from openpyxl import load_workbook

BOOK = "eToro-suitability-test-calculator.xlsx"
MON = f"'[{BOOK}]MONITORING'!"

CYSEC = "CySEC v24"
FCA = "FCA v15"
ASIC = "ASIC v9"
GAML = "ASIC GAML v15"
FSRA = "FSRA v10"
SHEETS = [CYSEC, FCA, ASIC, GAML, FSRA]

# config key -> the question id printed in column A of a calculator tab
QUESTION_IDS = {
    "Equities": "Q33", "Crypto": "Q34", "LeveragedCfd": "Q35",
    "EquitiesInvestedAmount": "Q47", "CryptoInvestedAmount": "Q48",
    "LeveragedCfdInvestedAmount": "Q45", "TradingKnowledge": "Q3",
    "TradingStrategy": "Q5", "TradingPurpose": "Q8", "RiskAppetite": "Q9",
    "IncomeSource": "Q15", "AnnualIncome": "Q10", "LiquidAssets": "Q11",
}
KA_NAMES = ["NewerLeverage", "NewerMarginCall", "NewerCfd",
            "NewerStopLossTrigger", "NewerStopLossGapThrough", "NewerCfdTrs"]
MICA_NAMES = ["MiCACryptoAssessmentHighVolatility", "MiCACryptoAssessmentCyberRisks",
              "MiCACryptoAssessmentRecoverLoss", "MiCACryptoAssessmentInvestingRisks",
              "MiCACryptoAssessmentPrivateKey"]
MICA_WRONG = ["Crypto makes a stable investment", "Using the same password everywhere",
              "Covers the recovery", "Minimal", "A shared password"]
# hard-block condition rows, keyed by the question id that opens their column-B label
BLOCK_CONDITIONS = ["Q9", "Q8", "Q10", "Q11", "Q15"]


def locate(sheet):
    """Map logical names onto cell addresses by scanning one tab's labels."""
    ws = load_workbook(BOOK)[sheet]
    col_a, col_b, col_f = {}, {}, {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=6):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            {1: col_a, 2: col_b, 6: col_f}.get(cell.column, {})[cell.value] = cell.row

    def by_prefix(prefix, column, required=True):
        for label, row in col_b.items():
            if label.startswith(prefix):
                return f"{column}{row}"
        if required:
            raise KeyError(f"{sheet}: no column-B label starting {prefix!r}")
        return None

    out = {}
    for i in range(1, 10):
        cell = by_prefix(f"Component {i} -", "E", required=i != 9)
        # component 9 only exists on the CySEC tab, and there its label carries a formula
        if cell and (i != 9 or sheet == CYSEC):
            out[f"C{i}"] = cell
    out.update({
        "KaTotal": by_prefix("Total assessment score", "E"),
        "FactorA": by_prefix("Factor A", "E"),
        "FactorB": by_prefix("Factor B", "E"),
        "LevelWeight": by_prefix("ClientRiskLevel", "E"),
        "LevelName": by_prefix("ClientRiskLevel", "D"),
        "AuthorizedRiskScore": by_prefix("AuthorizedRiskScore", "E"),
        "HardBlock": by_prefix("SuitabilityBlock", "E"),
        "Outcome": by_prefix("OUTCOME", "C"),
        "TargetVerdict": by_prefix("Can this user copy it?", "C"),
    })
    for qid in BLOCK_CONDITIONS:
        cell = by_prefix(f"{qid} ", "E", required=False)
        if cell:
            out[f"Cond{qid}"] = cell

    return {
        "inputs": {k: f"C{col_a[qid]}" for k, qid in QUESTION_IDS.items()},
        "ka": [f"C{col_f[n]}" for n in KA_NAMES],
        "mica": [f"C{col_f[n]}" for n in MICA_NAMES if n in col_f],
        "out": out,
    }


# Yes/No per statement, in KA_NAMES order. The two +2 statements are true, the rest false.
ALL_CORRECT = ["Yes", "Yes", "No", "No", "No", "No"]
NONE_TICKED = ["No"] * 6


def ticks(*names):
    """Build a Yes/No vector in KA_NAMES order from the answers the user ticked."""
    return ["Yes" if n in names else "No" for n in KA_NAMES]


# Every vector from WeightQuestionAnswerCalculatorTests.AnswerTestData in the KYCAnalyzer
# repo, with the RiskLevel the calculator asserts. These pin component 8 to the service.
ENGINE_VECTORS = [
    (ticks("NewerLeverage", "NewerCfd", "NewerMarginCall",
           "NewerStopLossTrigger", "NewerStopLossGapThrough"), 0, 2, "Medium"),
    (ticks("NewerLeverage", "NewerCfd", "NewerMarginCall",
           "NewerStopLossTrigger"), 4, 3, "MediumHigh"),
    (ticks("NewerLeverage", "NewerCfd", "NewerMarginCall"), 8, 4, "High"),
    (ticks("NewerLeverage", "NewerCfd"), 4, 3, "MediumHigh"),
    (ticks("NewerLeverage"), 8, 4, "High"),
    (ticks("NewerMarginCall"), 8, 4, "High"),
    (ticks("NewerLeverage", "NewerMarginCall"), 12, 4, "High"),
    (ticks("NewerLeverage", "NewerStopLossTrigger"), 4, 3, "MediumHigh"),
    (ticks("NewerLeverage", "NewerStopLossTrigger",
           "NewerStopLossGapThrough"), 0, 2, "Medium"),
    (ticks("NewerCfd", "NewerStopLossTrigger",
           "NewerStopLossGapThrough"), -8, 1, "Low"),
]
EXPERIENCED = {
    "Equities": "Above 20 times", "Crypto": "Above 20 times",
    "LeveragedCfd": "Above 40 times", "EquitiesInvestedAmount": "Above $2,000",
    "CryptoInvestedAmount": "Above $2,000", "LeveragedCfdInvestedAmount": "Above $2,000",
    "TradingKnowledge": "Professional certificate / work experience",
    "TradingStrategy": "A few seconds up to 24 hours",
    "TradingPurpose": "Short-term returns", "RiskAppetite": "80% / -48%",
    "IncomeSource": "Salary",
}
# The profile that trips the four common conditions: cautious saver, low income, low assets.
BLOCK_PROFILE = {"RiskAppetite": "5% / -3%", "TradingPurpose": "Saving for a home",
                 "AnnualIncome": "Up to $10K", "LiquidAssets": "$10K - $50K"}

CASES = [
    {
        "name": "shipped defaults - plausible mid-range profile under CySEC",
        "expect": {"C1": 2, "C2": 2, "C3": 2, "C4": 3, "C5": 3, "C6": 3, "C7": 4, "C8": 3,
                   "C9": 4, "KaTotal": 4, "FactorA": 3, "FactorB": 2,
                   "LevelName": "Medium", "AuthorizedRiskScore": 6, "HardBlock": "Not blocked"},
    },
    {
        "name": "a failed MiCA quiz drags an otherwise maxed-out CySEC user down a band",
        "inputs": EXPERIENCED,
        "ka": ALL_CORRECT,
        "mica": MICA_WRONG,
        "expect": {"C8": 4, "C9": 1, "FactorA": 4, "FactorB": 3,
                   "LevelName": "MediumHigh", "AuthorizedRiskScore": 8},
    },
    {
        "name": "nothing ticked on Q23 falls back to Medium, the score is discarded",
        "ka": NONE_TICKED,
        "expect": {"KaTotal": 4, "C8": 2},
    },
    {
        "name": "experienced short-term trader with a perfect quiz reaches the ceiling",
        "inputs": EXPERIENCED,
        "ka": ALL_CORRECT,
        "expect": {"KaTotal": 12, "C8": 4, "FactorA": 4, "FactorB": 4,
                   "LevelName": "High", "AuthorizedRiskScore": 10},
    },
    {
        "name": "Factor A caps an otherwise strong profile",
        "inputs": dict(EXPERIENCED, TradingPurpose="Future planning"),
        "expect": {"FactorA": 2, "LevelName": "Medium", "AuthorizedRiskScore": 6},
    },
    {
        "name": "pensioner with no experience who fails the quiz lands at the bottom",
        "inputs": {"IncomeSource": "Pension", "RiskAppetite": "10% / -6%",
                   "TradingPurpose": "Future planning",
                   "TradingKnowledge": "No financial knowledge",
                   "TradingStrategy": "More than several months / years"},
        "ka": ["No", "No", "Yes", "Yes", "Yes", "Yes"],
        "mica": MICA_WRONG,
        "expect": {"C7": 1, "C8": 1, "C9": 1, "FactorA": 2, "FactorB": 1,
                   "LevelName": "Low", "AuthorizedRiskScore": 3},
    },
]

# Component 9 exists only on the CySEC tab, so the same answers land a band higher
# everywhere else - Factor B averages six components instead of seven.
CASES += [
    {
        "name": f"{sheet} has no component 9, so the failed-MiCA user still reaches High",
        "sheet": sheet,
        "inputs": EXPERIENCED,
        "ka": ALL_CORRECT,
        "expect": {"FactorB": 4, "LevelName": "High", "AuthorizedRiskScore": 10},
    }
    for sheet in (FCA, ASIC, GAML, FSRA)
]

# Hard block. The four common conditions are identical in CySEC-24, FCA-15, ASIC-9 and
# FSRA-10; ASIC GAML-15 swaps two of them and adds a fifth on source of income. FCA-15
# is the only config that sets the check's DefaultResult to Blocked, and it makes no
# difference: BlockCalculator.Calculate never reads that field for a block with no
# nested checks, so FCA blocks exactly like CySEC.
CASES += [
    {
        "name": f"{sheet} blocks the cautious low-income saver",
        "sheet": sheet,
        "inputs": BLOCK_PROFILE,
        "expect": {"HardBlock": "BLOCKED", "CondQ9": "YES", "CondQ8": "YES",
                   "CondQ10": "YES", "CondQ11": "YES",
                   "Outcome": "Fully blocked from copy trading",
                   "TargetVerdict": "NO - fully blocked from copy"},
    }
    for sheet in (CYSEC, FCA, ASIC, FSRA)
]
CASES += [
    {
        "name": f"{sheet} does not block on low risk appetite alone",
        "sheet": sheet,
        "inputs": dict(BLOCK_PROFILE, AnnualIncome="$200K - $500K",
                       LiquidAssets="$200K - $500K"),
        "expect": {"HardBlock": "Not blocked", "CondQ9": "YES", "CondQ8": "YES",
                   "CondQ10": "no", "CondQ11": "no",
                   "LevelName": "Low", "AuthorizedRiskScore": 3},
    }
    for sheet in (CYSEC, FCA)
]
CASES += [
    {
        "name": "ASIC GAML needs a pension income source, so the common block profile passes",
        "sheet": GAML,
        "inputs": dict(BLOCK_PROFILE, IncomeSource="Salary"),
        "expect": {"HardBlock": "Not blocked", "CondQ8": "no", "CondQ15": "no"},
    },
    {
        "name": "ASIC GAML blocks a pensioner at 10% / -6%, which the others would not",
        "sheet": GAML,
        "inputs": {"RiskAppetite": "10% / -6%", "TradingPurpose": "Future planning",
                   "AnnualIncome": "Up to $10K", "LiquidAssets": "$10K - $50K",
                   "IncomeSource": "Pension"},
        "expect": {"HardBlock": "BLOCKED", "CondQ9": "YES", "CondQ15": "YES"},
    },
]
CASES += [
    {
        "name": f"the same pensioner profile is not blocked under {sheet}",
        "sheet": sheet,
        "inputs": {"RiskAppetite": "10% / -6%", "TradingPurpose": "Future planning",
                   "AnnualIncome": "Up to $10K", "LiquidAssets": "$10K - $50K",
                   "IncomeSource": "Pension"},
        "expect": {"HardBlock": "Not blocked", "CondQ9": "no"},
    }
    for sheet in (CYSEC, FCA, ASIC, FSRA)
]

# Q8 answers the tab's config does not list fall back to DefaultRiskLevel (Medium), not
# to the level CySEC assigns them. RiskAppetite is pinned high so Factor A = C5.
CASES += [
    {
        "name": "CySEC scores crypto-to-fiat High on Q8",
        "inputs": {"TradingPurpose": "Converting crypto to fiat",
                   "RiskAppetite": "80% / -48%"},
        "expect": {"C5": 4, "FactorA": 4},
    },
    {
        "name": "CySEC scores Investments Low, its most conservative Q8 answer",
        "inputs": {"TradingPurpose": "Investments", "RiskAppetite": "80% / -48%"},
        "expect": {"C5": 1, "FactorA": 1},
    },
    {
        "name": "no config scores answer 903, so it is Medium even under CySEC",
        "inputs": {"TradingPurpose": "Crypto trading and/or conversion",
                   "RiskAppetite": "80% / -48%"},
        "expect": {"C5": 2, "FactorA": 2},
    },
]
CASES += [
    {
        "name": f"{sheet} does not score crypto-to-fiat, so component 5 falls back to Medium",
        "sheet": sheet,
        "inputs": {"TradingPurpose": "Converting crypto to fiat",
                   "RiskAppetite": "80% / -48%"},
        "expect": {"C5": 2, "FactorA": 2},
    }
    for sheet in (FCA, ASIC, GAML, FSRA)
]
CASES += [
    {
        "name": f"{sheet} does not score Investments, so the fallback beats CySEC's Low",
        "sheet": sheet,
        "inputs": {"TradingPurpose": "Investments", "RiskAppetite": "80% / -48%"},
        "expect": {"C5": 2, "FactorA": 2},
    }
    for sheet in (FCA, FSRA)
]

# Statement 212 (NewerCfdTrs) is scored by every config but absent from the live question
# catalogue, so a real user always leaves it unticked and it always contributes +2. These
# cases pin the production-reachable ladder: totals -8, -4, 0, +4, +8, +12 and bands
# Low, Medium, Medium, MediumHigh, High, High. Correct judgement = tick the two true
# statements, leave the three answerable false ones alone.
_TRUE_ONES = ["NewerLeverage", "NewerMarginCall"]
_FALSE_ANSWERABLE = ["NewerCfd", "NewerStopLossTrigger", "NewerStopLossGapThrough"]

_PRODUCTION_LADDER = [
    # correct-of-5, ticked set, expected total, expected C8 band ordinal
    (0, _FALSE_ANSWERABLE, -8, 1),
    (1, ["NewerLeverage"] + _FALSE_ANSWERABLE, -4, 2),
    (2, _TRUE_ONES + _FALSE_ANSWERABLE, 0, 2),
    (3, ["NewerLeverage", "NewerCfd"], 4, 3),
    (4, ["NewerLeverage"], 8, 4),
    (5, _TRUE_ONES, 12, 4),
]

CASES += [
    {
        "name": (f"production ladder - {k} of 5 answerable statements correct "
                 f"-> total {total:+d}, C8 level {lvl}"),
        "ka": ["Yes" if n in ticked else "No" for n in KA_NAMES],
        "expect": {"KaTotal": total, "C8": lvl},
    }
    for k, ticked, total, lvl in _PRODUCTION_LADDER
]

CASES += [
    {
        "name": f"engine vector {i} - {sum(1 for t in vec if t == 'Yes')} ticked -> {level}",
        "ka": vec,
        "expect": {"KaTotal": total, "C8": weight},
    }
    for i, (vec, total, weight, level) in enumerate(ENGINE_VECTORS, start=1)
]


def unwrap(value):
    while hasattr(value, "value"):
        value = value.value
    if hasattr(value, "tolist"):
        flat = value.tolist()
        while isinstance(flat, list) and flat:
            flat = flat[0]
        return flat
    return value


def main():
    only = sys.argv[1] if len(sys.argv) > 1 else None
    if only and only not in SHEETS:
        print(f"unknown tab {only!r}; expected one of {SHEETS}")
        return 2

    maps = {s: locate(s) for s in SHEETS}
    model = formulas.ExcelModel().loads(BOOK).finish()
    failures = []
    ran = 0

    for case in CASES:
        sheet = case.get("sheet", CYSEC)
        if only and sheet != only:
            continue
        ran += 1
        ref = f"'[{BOOK}]{sheet.upper()}'!"
        m = maps[sheet]

        overrides = {ref + m["inputs"][k]: [[v]] for k, v in case.get("inputs", {}).items()}
        for cell, val in zip(m["ka"], case.get("ka", [])):
            overrides[ref + cell] = [[val]]
        for cell, val in zip(m["mica"], case.get("mica", [])):
            overrides[ref + cell] = [[val]]

        sol = model.calculate(inputs=overrides or None)
        got = {label: unwrap(sol[ref + cell]) for label, cell in m["out"].items()}

        comps = [got.get(f"C{i}", "-") for i in range(1, 10)]
        print(f"\n[{sheet}] {case['name']}")
        print(f"  C1..C9 = {comps}   quiz total = {got['KaTotal']}")
        print(f"  A={got['FactorA']}  B={got['FactorB']}  level={got['LevelName']}  "
              f"score={got['AuthorizedRiskScore']}  {got['HardBlock']}")

        for label, want in case["expect"].items():
            if label not in got:
                failures.append(f"[{sheet}] {case['name']}: {label} not present on this tab")
                continue
            actual = got[label]
            if str(actual).strip() != str(want).strip() and actual != want:
                failures.append(
                    f"[{sheet}] {case['name']}: {label} expected {want!r}, got {actual!r}")

    if not only:
        mon = load_workbook(BOOK)["Monitoring"]
        mcol = {c.value: c.row for r in mon.iter_rows(max_col=2) for c in r
                if isinstance(c.value, str)}
        sol = model.calculate()
        print()
        for label, cell in {"FSUST": f"D{mcol['FSUST']}",
                            "CU": f"D{mcol['CU (copy utilisation)']}",
                            "Outcome": f"C{mcol['OUTCOME']}"}.items():
            print(f"Monitoring {label}: {unwrap(sol[MON + cell])}")

    print()
    if failures:
        for f in failures:
            print("FAIL", f)
        return 1
    print(f"all {ran} cases passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
